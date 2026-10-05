import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Paleta de Cores Executiva Nobre
    BG_COLOR = RGBColor(10, 14, 26)       # #0A0E1A (Azul Marinho Quase Preto)
    CARD_BG = RGBColor(18, 26, 46)        # #121A2E (Vidro Escuro / Card)
    GOLD = RGBColor(245, 158, 11)         # #F59E0B (Dourado Âmbar)
    GOLD_LIGHT = RGBColor(253, 230, 138)  # #FDE68A (Dourado Suave)
    WHITE = RGBColor(255, 255, 255)       # Branco
    SLATE = RGBColor(148, 163, 184)       # #94A3B8 (Texto Secundário)
    BORDER_COLOR = RGBColor(38, 52, 85)   # #263455
    CARD_HIGHLIGHT = RGBColor(28, 38, 68)

    # Resolução de Imagens (Prioriza public/apresentacao_imgs local)
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(script_dir)
    public_imgs = os.path.join(project_root, "public", "apresentacao_imgs")

    img_capa = os.path.join(public_imgs, "capa.jpg")
    img_pastoral = os.path.join(public_imgs, "painel_pastoral.jpg")
    img_gestao = os.path.join(public_imgs, "gestao_eventos.jpg")

    # Fallback para artefatos do brain se necessário
    brain_dir = r"C:\Users\lynx\.gemini\antigravity-ide\brain\1ad8aa03-3e06-4c3b-8441-16b1bfc0114b"
    if not os.path.exists(img_capa):
        img_capa = os.path.join(brain_dir, "horeb_apresentacao_capa_pt_1790726217722.jpg")
    if not os.path.exists(img_pastoral):
        img_pastoral = os.path.join(brain_dir, "painel_pastoral_pt_1790726242788.jpg")
    if not os.path.exists(img_gestao):
        img_gestao = os.path.join(brain_dir, "gestao_financeira_eventos_pt_1790726271005.jpg")

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="PLATAFORMA HOREB • APRESENTAÇÃO PASTORAL"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.5), Inches(0.35))
        ctf = cat_box.text_frame
        ctf.word_wrap = True
        cp = ctf.paragraphs[0]
        cp.text = category_text.upper()
        cp.font.name = "Segoe UI"
        cp.font.size = Pt(11)
        cp.font.bold = True
        cp.font.color.rgb = GOLD

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.8))
        ttf = title_box.text_frame
        ttf.word_wrap = True
        tp = ttf.paragraphs[0]
        tp.text = title_text
        tp.font.name = "Segoe UI"
        tp.font.size = Pt(25)
        tp.font.bold = True
        tp.font.color.rgb = WHITE

    # =========================================================================
    # SLIDE 1: CAPA
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_background(s1)

    if os.path.exists(img_capa):
        s1.shapes.add_picture(img_capa, Inches(6.4), Inches(1.1), Inches(6.2), Inches(5.3))

    t1_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(5.3), Inches(5.3))
    tf1 = t1_box.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "PLATAFORMA ECLESIÁSTICA HOREB"
    p.font.name = "Segoe UI"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GOLD

    p = tf1.add_paragraph()
    p.text = "A Tecnologia a Serviço do Reino"
    p.font.name = "Segoe UI"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(14)

    p = tf1.add_paragraph()
    p.text = "Gestão ministerial de excelência aliada ao acolhimento contínuo de almas na era digital."
    p.font.name = "Segoe UI"
    p.font.size = Pt(16)
    p.font.color.rgb = SLATE
    p.space_after = Pt(22)

    p = tf1.add_paragraph()
    p.text = "“A tecnologia não substitui o pastor; ela capacita o ministério a alcançar a ovelha que está sofrendo em silêncio.”"
    p.font.name = "Segoe UI"
    p.font.size = Pt(13)
    p.font.italic = True
    p.font.color.rgb = GOLD_LIGHT
    p.space_after = Pt(18)

    p = tf1.add_paragraph()
    p.text = "Apresentação Oficial para Pastores Presidentes, Bispos e Diretorias de Igrejas"
    p.font.name = "Segoe UI"
    p.font.size = Pt(11)
    p.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 2: O GRANDE DILEMA DO MINISTÉRIO PASTORAL
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_background(s2)
    add_header(s2, "O Diagnóstico: O Abandono Silencioso e a Sobrecarga Pastoral")

    c_top = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.75), Inches(11.73), Inches(1.3))
    c_top.fill.solid()
    c_top.fill.fore_color.rgb = CARD_HIGHLIGHT
    c_top.line.color.rgb = GOLD
    c_top.line.width = Pt(1.5)
    ctf = c_top.text_frame
    ctf.word_wrap = True
    cp = ctf.paragraphs[0]
    cp.text = "O DESAFIO DA 'PORTA DOS FUNDOS' NA IGREJA:"
    cp.font.name = "Segoe UI"
    cp.font.size = Pt(12)
    cp.font.bold = True
    cp.font.color.rgb = GOLD
    cp2 = ctf.add_paragraph()
    cp2.text = "Mais de 40% das pessoas que se afastam da igreja não saem por rebeldia, mas por solidão, crise emocional não notada ou ausência de acolhimento no momento mais agudo da dor."
    cp2.font.name = "Segoe UI"
    cp2.font.size = Pt(14)
    cp2.font.color.rgb = WHITE

    cards_data = [
        ("Sobrecarga da Liderança", "O pastor titular e sua equipe têm coração pastoral, mas o dia só tem 24h. Centenas de mensagens se acumulam sem resposta humanizada a tempo."),
        ("Comunicação Caótica", "Avisos espirituais e pedidos urgentes de oração se perdem no ruído de grupos paralelos de WhatsApp, sem histórico e sem sigilo ético."),
        ("Ausência de Indicadores", "Dificuldade em acompanhar a frequência das células nos lares, tarefas de departamentos e emitir comprovantes confiáveis de dízimos para IRPF.")
    ]

    for i, (title, desc) in enumerate(cards_data):
        left_pos = Inches(0.8 + i * 4.0)
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(3.35), Inches(3.73), Inches(3.4))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = BORDER_COLOR
        c.line.width = Pt(1)
        tf = c.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"0{i+1}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = GOLD
        p.space_after = Pt(6)

        p = tf.add_paragraph()
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.space_after = Pt(10)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = "Segoe UI"
        p.font.size = Pt(12)
        p.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 3: O QUE É O HOREB? ECOSSISTEMA MINISTERIAL INTEGRADO
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_background(s3)
    add_header(s3, "O Que é o Horeb: Um Ecossistema Ministerial Integrado")

    pilares = [
        ("01. Cuidado de Almas 24/7", 
         "Aconselhamento bíblico com Inteligência Artificial humanizada com voz que escuta, ora pelo membro e identifica crises emocionais em tempo real."),
        ("02. Comunhão, Departamentos e Células", 
         "Workspaces de ministérios com atas e tarefas, localização de células nos lares, devocionais diários com áudio e convite nominal de amigos."),
        ("03. Governança, Eventos e Transparência", 
         "Gestão financeira com dízimos nominais para IRPF, bilheteria com QR Code e envio automático de ingressos por e-mail, e controle multi-tenant de congregações.")
    ]

    for i, (title, desc) in enumerate(pilares):
        top_pos = Inches(1.8 + i * 1.75)
        c = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.73), Inches(1.5))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = GOLD if i == 0 else BORDER_COLOR
        c.line.width = Pt(1.5 if i == 0 else 1)
        tf = c.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = GOLD if i == 0 else WHITE
        p.space_after = Pt(6)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = "Segoe UI"
        p.font.size = Pt(13)
        p.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 4: CHECK-IN DE ALMA & IA PASTORAL COM ÁUDIO
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_background(s4)
    add_header(s4, "Check-in de Alma: Acolhimento Humanizado com Reprodução em Áudio")

    if os.path.exists(img_pastoral):
        s4.shapes.add_picture(img_pastoral, Inches(6.0), Inches(1.7), Inches(6.5), Inches(5.0))

    t4_box = s4.shapes.add_textbox(Inches(0.8), Inches(1.65), Inches(4.9), Inches(5.1))
    tf4 = t4_box.text_frame
    tf4.word_wrap = True

    pontos_s4 = [
        ("Termômetro Emocional Bíblico:", "O membro escolhe como se sente (Feliz, Grato, Cansado, Ansioso, Triste) com design respeitoso e acolhedor."),
        ("Tratamento Nominal Afetuoso:", "A IA reconhece o primeiro nome do membro (ex: 'Graça e Paz, irmão Luan!'), oferecendo leitura sensível e pastoral."),
        ("Player de Áudio Nativo (Voz Brasileira):", "O usuário pode escutar a meditação e a oração geradas em voz natural fluida a qualquer hora do dia ou da noite."),
        ("100% Fiel à Doutrina Cristã:", "Respostas estritamente embasadas nas Escrituras Sagradas, confortando sem frieza robótica ou conselhos seculares vazios.")
    ]

    for i, (head, body) in enumerate(pontos_s4):
        p = tf4.add_paragraph() if i > 0 else tf4.paragraphs[0]
        p.text = head
        p.font.name = "Segoe UI"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = GOLD
        p.space_after = Pt(2)

        p2 = tf4.add_paragraph()
        p2.text = body
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(12)
        p2.font.color.rgb = SLATE
        p2.space_after = Pt(10)

    # =========================================================================
    # SLIDE 5: DA IA AO PASTOR HUMANO - A PONTE SEGURA
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_background(s5)
    add_header(s5, "A Ponte do Cuidado: Da Inteligência Artificial ao Pastor Humano")

    steps_s5 = [
        ("Passo 1: Acolhimento Diário 24/7", "A IA atua como primeiro apoio contínuo, consolando e intercedendo pelas situações rotineiras com a Palavra."),
        ("Passo 2: Detecção Inteligente de Crise", "Em casos de luto, depressão severa ou conflito agudo, o sistema aciona um botão de alerta: 'Falar com o Pastor'."),
        ("Passo 3: Fila ao Vivo ou Agendamento", "Se houver pastor online no painel, o membro é atendido na hora. Caso contrário, escolhe data e horário na agenda ministerial."),
        ("Passo 4: Preparo Espiritual Prévio", "O pastor recebe o histórico e resumo prévio do motivo, entrando no gabinete presencial ou virtual já preparado em oração.")
    ]

    for i, (title, desc) in enumerate(steps_s5):
        left_pos = Inches(0.8 + (i % 2) * 6.0)
        top_pos = Inches(1.85 + (i // 2) * 2.5)
        c = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, Inches(5.7), Inches(2.2))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = GOLD if i == 1 else BORDER_COLOR
        c.line.width = Pt(1.5 if i == 1 else 1)
        tf = c.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = GOLD if i == 1 else WHITE
        p.space_after = Pt(6)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = "Segoe UI"
        p.font.size = Pt(13)
        p.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 6: GESTÃO DE MINISTÉRIOS E DEPARTAMENTOS (WORKSPACES) [NOVO!]
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_background(s6)
    add_header(s6, "Gestão de Ministérios & Departamentos: Workspaces Exclusivos", "DEPARTAMENTOS ECLESIAIS • GOVERNANÇA ATIVA")

    min_cards = [
        ("Workspaces por Ministério", "Espaço independente para Louvor, Jovens, Mulheres, Kids, Homens, Ação Social e Comunicação, com identidade própria."),
        ("Atas de Reuniões Oficiais", "Registro digital de atas, decisões, aprovações e pautas das reuniões ministeriais, arquivadas em nuvem com histórico auditável."),
        ("Gestão de Tarefas & Escalas", "Distribuição clara de atividades, prazos e responsabilidades entre voluntários e obreiros, sem desencontro de informações."),
        ("Equipe com Fotos e Lideranças", "Visualização dos membros do departamento com fotos de perfil, telefones e funções bem definidas.")
    ]

    for i, (title, desc) in enumerate(min_cards):
        left_pos = Inches(0.8 + (i % 2) * 6.0)
        top_pos = Inches(1.85 + (i // 2) * 2.5)
        c = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, Inches(5.7), Inches(2.2))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = GOLD if i == 0 else BORDER_COLOR
        c.line.width = Pt(1.5 if i == 0 else 1)
        tf = c.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = GOLD if i == 0 else WHITE
        p.space_after = Pt(6)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = "Segoe UI"
        p.font.size = Pt(13)
        p.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 7: CÉLULAS E DISCIPULADO NOS LARES
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_background(s7)
    add_header(s7, "Células e Discipulado: A Igreja Conectada nos Lares")

    cel_data = [
        ("Mapa da Igreja nos Bairros", "Visitantes e novos decididos encontram em 1 clique a célula mais próxima de sua casa, com endereço, dia, horário e contato do líder."),
        ("Alinhamento Doutrinário Oficial", "Os líderes de célula acessam diretamente no aplicativo os estudos e esboços oficiais enviados pela liderança pastoral da igreja."),
        ("Frequência e Acompanhamento", "O líder registra os membros presentes na reunião semanal, permitindo que a liderança pastoral saiba quem está faltando antes que esfrie na fé.")
    ]

    for i, (title, desc) in enumerate(cel_data):
        top_pos = Inches(1.85 + i * 1.65)
        c = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.73), Inches(1.4))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = BORDER_COLOR
        c.line.width = Pt(1)
        tf = c.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = GOLD
        p.space_after = Pt(4)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = "Segoe UI"
        p.font.size = Pt(13)
        p.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 8: MULTIPLICAÇÃO & EVANGELISMO - "CONVIDAR AMIGO" [NOVO!]
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_background(s8)
    add_header(s8, "Multiplicação e Evangelismo: 'Convidar Amigo' em 1 Clique", "CRESCIMENTO DA IGREJA • ENGAJAMENTO")

    convite_cards = [
        ("Convite Fácil via WhatsApp & E-mail", "Qualquer membro pode convidar vizinhos, amigos e familiares gerando um link acolhedor pré-preenchido com o nome do convidado."),
        ("Página de Boas-Vindas Desbloqueada", "O amigo acessa uma tela limpa e elegante na web (/[slug]/cadastro) com a identidade visual e o logo da igreja, sem complicação."),
        ("Blindagem Restrita de Membresia", "O sistema garante que novos cadastros convidados por membros recebam estritamente o papel de Membro, sem risco de acesso indevido."),
        ("Acompanhamento pela Secretaria", "A liderança e a secretaria visualizam no painel quem convidou, quem aceitou e quando o novo membro concluiu seu cadastro.")
    ]

    for i, (title, desc) in enumerate(convite_cards):
        left_pos = Inches(0.8 + (i % 2) * 6.0)
        top_pos = Inches(1.85 + (i // 2) * 2.5)
        c = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, Inches(5.7), Inches(2.2))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = GOLD if i == 0 else BORDER_COLOR
        c.line.width = Pt(1.5 if i == 0 else 1)
        tf = c.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = GOLD if i == 0 else WHITE
        p.space_after = Pt(6)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = "Segoe UI"
        p.font.size = Pt(13)
        p.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 9: GESTÃO FINANCEIRA E PRESTAÇÃO DE CONTAS (IRPF)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_background(s9)
    add_header(s9, "Gestão Financeira Transparente & Comprovantes de Dízimo para IRPF")

    if os.path.exists(img_gestao):
        s9.shapes.add_picture(img_gestao, Inches(5.8), Inches(1.7), Inches(6.7), Inches(5.0))

    t9_box = s9.shapes.add_textbox(Inches(0.8), Inches(1.65), Inches(4.7), Inches(5.1))
    tf9 = t9_box.text_frame
    tf9.word_wrap = True

    pontos_s9 = [
        ("Dízimos e Ofertas Nominais:", "Rastreamento seguro via PIX, cartão e transferências bancárias com identificação individual do membro."),
        ("Comprovante Anual para IRPF:", "Gera extrato oficial com CNPJ da igreja para a declaração de Imposto de Renda dos membros de forma imediata."),
        ("Transparência com a Diretoria:", "Relatórios mensais categorizados (Missões, Aluguel, Reforma, Obras Sociais) prontos para conselho fiscal e assembleia."),
        ("Segurança Bancária:", "Sem manipulação manual desprotegida de dinheiro em espécie na tesouraria.")
    ]

    for i, (head, body) in enumerate(pontos_s9):
        p = tf9.add_paragraph() if i > 0 else tf9.paragraphs[0]
        p.text = head
        p.font.name = "Segoe UI"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = GOLD
        p.space_after = Pt(2)

        p2 = tf9.add_paragraph()
        p2.text = body
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(12)
        p2.font.color.rgb = SLATE
        p2.space_after = Pt(10)

    # =========================================================================
    # SLIDE 10: EVENTOS, BILHETERIA & QR CODE POR E-MAIL [ATUALIZADO!]
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_background(s10)
    add_header(s10, "Eventos e Congressos: Bilheteria Digital, E-mail com QR Code & Portaria")

    ev_cards = [
        ("Inscrições e Vendas Sem Taxas", "Congressos de mulheres, retiros de jovens e conferências ministeriais (pagos ou gratuitos) com inscrição direta no aplicativo."),
        ("Disparo Automático por E-mail", "Ao confirmar o ingresso, o participante recebe e-mail oficial com o QR Code individualizado e código único de validação."),
        ("Catraca Scanner em Segundos", "Os obreiros e voluntários realizam a leitura do QR Code na entrada usando a própria câmera do celular, eliminando filas e tumultos."),
        ("Carteira Digital do Membro", "Todos os ingressos adquiridos ficam salvos na carteira digital do usuário no app para exibição offline a qualquer momento.")
    ]

    for i, (title, desc) in enumerate(ev_cards):
        left_pos = Inches(0.8 + (i % 2) * 6.0)
        top_pos = Inches(1.85 + (i // 2) * 2.5)
        c = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, Inches(5.7), Inches(2.2))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = GOLD if i == 1 else BORDER_COLOR
        c.line.width = Pt(1.5 if i == 1 else 1)
        tf = c.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = GOLD if i == 1 else WHITE
        p.space_after = Pt(6)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = "Segoe UI"
        p.font.size = Pt(13)
        p.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 11: GOVERNANÇA EM REDE: SEDES, FILIAIS & RBAC MULTI-TENANT
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_background(s11)
    add_header(s11, "Governança em Rede: Sedes, Filiais e Segurança Multi-Tenant (RBAC)")

    roles = [
        ("Pastor Presidente / Bispo (Superadmin)", "Visão panorâmica consolidada de toda a rede de congregações, relatórios globais de engajamento e troca ágil entre filiais."),
        ("Pastores de Congregações Locais", "Autonomia na gestão dos membros, aconselhamentos, células e eventos da sua congregação filial."),
        ("Tesouraria & Conselho Fiscal", "Acesso restrito aos módulos financeiros e prestação de contas, sem acesso a dados sigilosos pastorais."),
        ("Líderes de Ministérios & Membros", "Gestão focada do seu próprio departamento, com membros acessando devocionais, células e seus próprios dízimos.")
    ]

    for i, (role, desc) in enumerate(roles):
        left_pos = Inches(0.8 + (i % 2) * 6.0)
        top_pos = Inches(1.85 + (i // 2) * 2.5)
        c = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, Inches(5.7), Inches(2.2))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = GOLD if i == 0 else BORDER_COLOR
        c.line.width = Pt(1.5 if i == 0 else 1)
        tf = c.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = role
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = GOLD if i == 0 else WHITE
        p.space_after = Pt(6)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = "Segoe UI"
        p.font.size = Pt(13)
        p.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 12: IMPLANTAÇÃO SEM TRAVAMENTO EM 7 DIAS
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_background(s12)
    add_header(s12, "Cronograma de Implantação Rápida em 7 Dias")

    dias = [
        ("Dias 1 e 2", "Configuração Inicial", "Inclusão da logo da congregação, cores institucionais, filiais e escala de pastores."),
        ("Dias 3 e 4", "Treinamento Pastoral", "Capacitação prática da secretaria, tesouraria e equipe pastoral no painel."),
        ("Dias 5 e 6", "Células e Ministérios", "Inserção dos pequenos grupos, criação dos workspaces ministeriais e agenda."),
        ("Dia 7", "Lançamento no Culto", "Apresentação com QR Code no telão do templo para toda a igreja baixar o app.")
    ]

    for i, (dia, title, desc) in enumerate(dias):
        left_pos = Inches(0.8 + i * 2.98)
        c = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(1.95), Inches(2.78), Inches(4.7))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = GOLD if i == 3 else BORDER_COLOR
        c.line.width = Pt(1.5 if i == 3 else 1)
        tf = c.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = dia
        p.font.name = "Segoe UI"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = GOLD
        p.space_after = Pt(4)

        p = tf.add_paragraph()
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.space_after = Pt(12)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = "Segoe UI"
        p.font.size = Pt(12)
        p.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 13: CONCLUSÃO E PRÓXIMO PASSO
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_background(s13)

    c13 = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.2), Inches(10.33), Inches(5.1))
    c13.fill.solid()
    c13.fill.fore_color.rgb = CARD_BG
    c13.line.color.rgb = GOLD
    c13.line.width = Pt(2)
    tf13 = c13.text_frame
    tf13.word_wrap = True

    p = tf13.paragraphs[0]
    p.text = "A DECISÃO ESTRATÉGICA DO MINISTÉRIO"
    p.font.name = "Segoe UI"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GOLD
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(14)

    p = tf13.add_paragraph()
    p.text = "“Investir no Horeb não é contratar um sistema;\né fechar a porta dos fundos da igreja e cuidar com excelência\nde cada vida que Deus confiou ao seu ministério.”"
    p.font.name = "Segoe UI"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(26)

    p = tf13.add_paragraph()
    p.text = "Próximo Passo:"
    p.font.name = "Segoe UI"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = GOLD
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(6)

    p = tf13.add_paragraph()
    p.text = "Agende uma demonstração prática de 20 minutos com a equipe pastoral e diretoria da sua igreja.\nO Horeb está pronto para transformar o cuidado com as suas ovelhas."
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.color.rgb = SLATE
    p.alignment = PP_ALIGN.CENTER

    # Salva na raiz e em public/ para download direto
    root_output = os.path.join(project_root, "Apresentacao_Horeb_Pastores.pptx")
    public_output = os.path.join(project_root, "public", "Apresentacao_Horeb_Pastores.pptx")
    
    prs.save(root_output)
    print(f"Salvo em: {root_output}")
    prs.save(public_output)
    print(f"Salvo em: {public_output}")

if __name__ == "__main__":
    create_deck()
